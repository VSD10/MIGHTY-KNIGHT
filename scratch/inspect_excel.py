import openpyxl
import os

path = os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
path = os.path.abspath(path)
print('Path:', path)
print('Exists:', os.path.exists(path))

wb = openpyxl.load_workbook(path)
print('Sheets:', wb.sheetnames)
ws = wb.active
print('Sheet name:', ws.title)
print('Max row:', ws.max_row)
print('Max col:', ws.max_column)
print()

print('--- Row 1 (all non-empty cols) ---')
for col in range(1, ws.max_column + 1):
    v = ws.cell(1, col).value
    if v is not None:
        print(f'  Col {col}: {repr(v)}')

print()
print('--- Row 2 (all non-empty cols) ---')
for col in range(1, ws.max_column + 1):
    v = ws.cell(2, col).value
    if v is not None:
        print(f'  Col {col}: {repr(v)}')

print()
print('--- Row 3 (all non-empty cols) ---')
for col in range(1, ws.max_column + 1):
    v = ws.cell(3, col).value
    if v is not None:
        print(f'  Col {col}: {repr(v)}')

print()
print('--- First 5 student rows (rows 4-8), first 15 cols ---')
for row in range(4, 9):
    print(f'  Row {row}:')
    for col in range(1, 16):
        v = ws.cell(row, col).value
        if v is not None:
            print(f'    Col {col}: {repr(v)}')

print()
print('--- Sample date columns (cols 7-20) for row 4 ---')
for col in range(7, 21):
    v = ws.cell(4, col).value
    h1 = ws.cell(1, col).value
    h2 = ws.cell(2, col).value
    print(f'  Col {col}: header1={repr(h1)}, header2={repr(h2)}, data={repr(v)}')

print()
print('--- Check merged cells ---')
for merge in list(ws.merged_cells.ranges)[:20]:
    print(f'  {merge}')
