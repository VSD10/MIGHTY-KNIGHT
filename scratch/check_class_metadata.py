import openpyxl
import os
import datetime
from collections import defaultdict

path = os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
wb = openpyxl.load_workbook(os.path.abspath(path), data_only=True)
ws = wb.active

date_cols = {}
for col in range(8, ws.max_column, 2):
    d_val = ws.cell(3, col).value
    day_val = ws.cell(1, col).value
    if d_val:
        d_obj = d_val.date() if isinstance(d_val, datetime.datetime) else d_val
        date_cols[col] = (d_obj, str(day_val).strip() if day_val else d_obj.strftime("%A"), col + 1)

classes = defaultdict(list)

for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if not name:
        continue
    stud_id = ws.cell(r, 2).value or f"MKS{r:05d}"
    level = ws.cell(r, 4).value or "Basic"
    batch = ws.cell(r, 5).value or "General"
    
    for col, (d_obj, day_name, tr_col) in date_cols.items():
        t_val = ws.cell(r, col).value
        tr_val = ws.cell(r, tr_col).value
        if t_val is not None:
            t_str = t_val.strftime("%I:%M %p").lstrip('0') if isinstance(t_val, datetime.time) else str(t_val).strip()
            tr_str = str(tr_val).strip() if tr_val else "Unassigned"
            classes[(d_obj, t_str, tr_str)].append({
                "name": str(name).strip(),
                "id": str(stud_id).strip(),
                "level": str(level).strip(),
                "batch": str(batch).strip()
            })

print(f"Total unique classes: {len(classes)}")
multi_level = 0
multi_batch = 0
for k, stus in classes.items():
    lvls = set(s['level'] for s in stus)
    batches = set(s['batch'] for s in stus)
    if len(lvls) > 1:
        multi_level += 1
    if len(batches) > 1:
        multi_batch += 1

print(f"Classes with multiple levels: {multi_level}")
print(f"Classes with multiple batches: {multi_batch}")
