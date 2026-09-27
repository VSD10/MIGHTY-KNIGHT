import openpyxl
import os
import datetime
from collections import defaultdict, Counter

path = os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
wb = openpyxl.load_workbook(os.path.abspath(path), data_only=True)
ws = wb.active

date_cols = {}
for col in range(8, ws.max_column, 2):
    d_val = ws.cell(3, col).value
    day_val = ws.cell(1, col).value
    if d_val:
        d_obj = d_val.date() if isinstance(d_val, datetime.datetime) else d_val
        date_cols[col] = (d_obj, day_val, col + 1)

# Count how many of each weekday exist in October 2026
# Oct 2026 has 31 days:
# Oct 1 is Thursday
weekday_date_counts = Counter(day for d, day, tr in date_cols.values())
print("October 2026 weekday occurrences in calendar:")
for w, c in weekday_date_counts.items():
    print(f"  {w}: {c} times")

# Check for students:
# If student attends Monday, Wednesday, Friday:
# There are 4 Mondays, 4 Wednesdays, 5 Fridays in October = 13 possible days.
# But Planned is 12! Which Friday was skipped?
skipped_details = []
for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if not name:
        continue
    planned = ws.cell(r, 6).value
    
    assigned_dates = []
    weekday_slots = {}
    for col, (d_obj, day_val, tr_col) in date_cols.items():
        t_val = ws.cell(r, col).value
        tr_val = ws.cell(r, tr_col).value
        if t_val is not None:
            assigned_dates.append(d_obj)
            weekday_slots[day_val] = (t_val, tr_val)
            
    num_assigned = len(assigned_dates)
    
    # What dates would match the student's recurring weekdays in October?
    matching_calendar_dates = [d_obj for col, (d_obj, day_val, tr_col) in date_cols.items() if day_val in weekday_slots]
    
    if len(matching_calendar_dates) != num_assigned:
        skipped = set(matching_calendar_dates) - set(assigned_dates)
        skipped_details.append((name, planned, num_assigned, len(matching_calendar_dates), sorted(skipped)))

print(f"\nStudents where assigned dates count != all calendar occurrences of their weekdays: {len(skipped_details)}")
print("Sample skipped dates:")
for item in skipped_details[:10]:
    print(f"  {item[0]}: planned={item[1]}, assigned={item[2]}, all_weekday_dates={item[3]}, skipped={[d.strftime('%Y-%m-%d') for d in item[4]]}")
