import openpyxl
from collections import defaultdict
import re

wb = openpyxl.load_workbook('sample_data/Oct26_Schedule_Final.xlsx', data_only=True)
ws = wb.active
students = []
for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 1).value
    sid = ws.cell(r, 2).value
    raw_b = ws.cell(r, 5).value
    if name and sid:
        students.append({'name': str(name).strip(), 'id': str(sid).strip(), 'raw_b': str(raw_b or '').strip()})

by_raw = defaultdict(list)
for s in students:
    by_raw[s['raw_b']].append(s)

for b, slist in sorted(by_raw.items()):
    names = [s['name'] for s in slist[:4]]
    print(f'{b:28} -> {len(slist):2d} students: {names}')
