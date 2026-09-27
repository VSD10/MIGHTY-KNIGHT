import openpyxl
import os
import datetime

path = os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
wb = openpyxl.load_workbook(os.path.abspath(path), data_only=True)
ws = wb.active

sunday_cols = []
for col in range(8, ws.max_column, 2):
    day_val = ws.cell(1, col).value
    if day_val and str(day_val).strip().lower() == "sunday":
        d_val = ws.cell(3, col).value
        sunday_cols.append((col, d_val))

print(f"Sunday columns: {len(sunday_cols)}")

sunday_times = set()
for r in range(4, ws.max_row + 1):
    for col, d_val in sunday_cols:
        t_val = ws.cell(r, col).value
        tr_val = ws.cell(r, col + 1).value
        if t_val is not None:
            sunday_times.add((str(t_val), str(tr_val)))

print("Sunday times and trainers in Oct 26 Excel:")
for t, tr in sorted(sunday_times):
    print(f"  {t} with {tr}")
