import openpyxl
import os
import datetime
from collections import defaultdict

path = os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
wb = openpyxl.load_workbook(os.path.abspath(path), data_only=True)
ws = wb.active

date_cols = {}
for col in range(8, ws.max_column, 2):
    d_val = ws.cell(3, col).value
    if d_val:
        d_obj = d_val.date() if isinstance(d_val, datetime.datetime) else d_val
        date_cols[col] = d_obj

student_date_duplicates = []
for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if not name:
        continue
    assigned_dates = []
    for col, d_obj in date_cols.items():
        t_val = ws.cell(r, col).value
        if t_val is not None:
            assigned_dates.append(d_obj)
    if len(assigned_dates) != len(set(assigned_dates)):
        student_date_duplicates.append(name)

print(f"Students with duplicate classes on same date: {len(student_date_duplicates)}")
