import openpyxl
from collections import defaultdict
wb = openpyxl.load_workbook(r"sample_data/Oct'26 Schedule_FRESH-1.xlsx", data_only=True)
ws = wb['OCTOBER 2026 SCHEDULE']
day_coaches = defaultdict(lambda: defaultdict(int))
coach_class_counts = defaultdict(int)
for r in range(4, 131):
    sid = ws.cell(r, 2).value
    for c in range(8, 70, 2):
        d_val = ws.cell(3, c).value
        if not d_val: continue
        d_str = d_val.strftime('%Y-%m-%d')
        t_val = ws.cell(r, c).value
        tr_val = ws.cell(r, c+1).value
        if t_val and tr_val:
            day_coaches[d_str][tr_val.strip()] += 1

print("First 10 days coach distribution in Excel:")
for d in sorted(day_coaches.keys())[:10]:
    print(d, dict(day_coaches[d]))
