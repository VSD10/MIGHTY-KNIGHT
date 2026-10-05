import os
import sys
import json
import sqlite3
import datetime
import re
from collections import defaultdict
import openpyxl

# Add backend directory to sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKEND_DIR = os.path.join(BASE_DIR, "backend")
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from app.models.student import StudentModel
from app.models.coach import CoachModel
from app.models.schedule import ScheduleResult, ScheduledClass, CoachCommunicationSlot
from app.utils.time_utils import parse_time_slot_sort_key
from app.storage.database import compute_master_data_fingerprint

EXCEL_PATH = os.path.join(BASE_DIR, "sample_data", "Oct26_Schedule_Final.xlsx")
DB_PATH = os.path.join(BACKEND_DIR, "data", "chess_scheduler.db")
JSON_PATH = os.path.join(BACKEND_DIR, "data", "reference_oct_schedule.json")

DAY_NAMES = {
    1: 'Thursday', 2: 'Friday', 3: 'Saturday', 4: 'Sunday', 5: 'Monday', 6: 'Tuesday', 7: 'Wednesday',
    8: 'Thursday', 9: 'Friday', 10: 'Saturday', 11: 'Sunday', 12: 'Monday', 13: 'Tuesday', 14: 'Wednesday',
    15: 'Thursday', 16: 'Friday', 17: 'Saturday', 18: 'Sunday', 19: 'Monday', 20: 'Tuesday', 21: 'Wednesday',
    22: 'Thursday', 23: 'Friday', 24: 'Saturday', 25: 'Sunday', 26: 'Monday', 27: 'Tuesday', 28: 'Wednesday',
    29: 'Thursday', 30: 'Friday', 31: 'Saturday'
}

LEVEL_RULES = [
    ('Basic 1', 1, r'basic\s*1'),
    ('Basic 2', 2, r'basic\s*2'),
    ('Beginner 1', 3, r'beginner\s*1'),
    ('Beginner 2', 4, r'beginner\s*2'),
    ('Early Intermediate 1', 5, r'early\s*int[e]?rmediate\s*1'),
    ('Early Intermediate 2', 6, r'early\s*int[e]?rmediate\s*2'),
    ('Intermediate 1', 7, r'(?<!early\s)int[e]?rmediate\s*1'),
    ('Intermediate 2', 8, r'(?<!early\s)int[e]?rmediate\s*2'),
    ('Intermediate 2', 8, r'intermediate$')
]

def map_batch_to_level(batch_str):
    b = str(batch_str or '').strip()
    btype = 'G'
    if b.startswith(('L', 'l')) or 'Limited' in b:
        btype = 'L'
    elif b.startswith(('I', 'i')) or 'Individual' in b:
        btype = 'I'
    for lvl_name, grp_no, pat in LEVEL_RULES:
        if re.search(pat, b, re.IGNORECASE):
            return lvl_name, grp_no, btype
    return 'Basic 1', 1, btype

def parse_and_feed():
    print(f"Loading authoritative workbook from: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['Oct26 Schedule corrections']

    students = []
    student_map = {}
    coaches_set = set()
    batches_map = {} # batch_id -> dict
    sessions = {}    # (date_str, time_slot, coach_name) -> session dict

    # 1. Parse Students & Class Slots
    for r in range(2, 124):
        sid = ws.cell(r, 2).value
        name = ws.cell(r, 1).value
        if not sid or r in [114, 115, 116, 117]:
            continue

        sid = str(sid).strip()
        name = str(name).strip()
        rating = ws.cell(r, 3).value
        raw_level = ws.cell(r, 4).value
        batch_raw = ws.cell(r, 5).value or 'G Basic1'
        batch_name = str(batch_raw).strip()
        avail_note = ws.cell(r, 6).value
        planned = ws.cell(r, 7).value or 12
        actual_cell = ws.cell(r, 8).value
        comments = ws.cell(r, 72).value

        norm_level, grp_no, btype = map_batch_to_level(batch_name)

        notes_parts = []
        if avail_note:
            notes_parts.append(str(avail_note).strip())
        if comments:
            notes_parts.append(str(comments).strip())
        combined_notes = ' | '.join(notes_parts)

        # Register batch
        batch_id = f"BATCH_{re.sub(r'[^a-zA-Z0-9]', '_', batch_name).upper()}"
        if batch_id not in batches_map:
            cap_min = 4 if btype == 'G' else 1
            cap_max = 8 if btype == 'G' else (3 if btype == 'L' else 1)
            batches_map[batch_id] = {
                "batch_id": batch_id,
                "batch_name": batch_name,
                "level": norm_level,
                "batch_type": btype,
                "group_number": grp_no,
                "capacity_min": cap_min,
                "capacity_max": cap_max,
                "fixed_trainer": "Unassigned",
                "schedule_timings": "",
                "student_ids": [],
                "students": []
            }

        batches_map[batch_id]["student_ids"].append(sid)
        batches_map[batch_id]["students"].append({
            "student_id": sid,
            "student_name": name,
            "student_level": norm_level,
            "mkca_rating": rating,
            "fixed_trainer": "Unassigned"
        })

        # Scheduled dates
        scheduled_dates = []
        recurring_slots = {} # day_name -> (time_slot, coach_name)

        for d in range(1, 32):
            tc = 10 + (d - 1) * 2
            cc = 11 + (d - 1) * 2
            t_val = ws.cell(r, tc).value
            c_val = ws.cell(r, cc).value

            if t_val is not None or c_val is not None:
                scheduled_dates.append(d)
                if isinstance(t_val, (datetime.time, datetime.datetime)):
                    start_hr = t_val.hour
                    start_mn = t_val.minute
                    end_hr = (start_hr + 1) % 24
                    end_mn = start_mn
                    s_ampm = 'AM' if start_hr < 12 else 'PM'
                    e_ampm = 'AM' if end_hr < 12 else 'PM'
                    s_h12 = start_hr if 1 <= start_hr <= 12 else (12 if start_hr in [0, 12] else start_hr - 12)
                    e_h12 = end_hr if 1 <= end_hr <= 12 else (12 if end_hr in [0, 12] else end_hr - 12)
                    time_slot = f"{s_h12:02d}:{start_mn:02d} {s_ampm} - {e_h12:02d}:{end_mn:02d} {e_ampm}"
                else:
                    time_slot = str(t_val).strip() if t_val else '06:00 PM - 07:00 PM'

                coach_name = str(c_val).strip() if c_val else 'Unassigned'
                coaches_set.add(coach_name)
                date_str = f"2026-10-{d:02d}"
                day_name = DAY_NAMES[d]

                if day_name not in recurring_slots:
                    recurring_slots[day_name] = (time_slot, coach_name)

                sess_key = (date_str, time_slot, coach_name)
                if sess_key not in sessions:
                    sessions[sess_key] = {
                        "date": date_str,
                        "day": day_name,
                        "time_slot": time_slot,
                        "coach_name": coach_name,
                        "student_ids": [],
                        "student_names": [],
                        "batches": set(),
                        "levels": set(),
                        "batch_types": set()
                    }
                sessions[sess_key]["student_ids"].append(sid)
                sessions[sess_key]["student_names"].append(name)
                sessions[sess_key]["batches"].add(batch_name)
                sessions[sess_key]["levels"].add(norm_level)
                sessions[sess_key]["batch_types"].add(btype)

        # Planned integer
        try:
            planned_cnt = int(planned)
        except Exception:
            planned_cnt = 12

        # Rating float
        try:
            rating_flt = float(rating) if rating is not None else None
        except Exception:
            rating_flt = None

        # Preferred days from scheduled days
        p_days = set(DAY_NAMES[d] for d in scheduled_dates)
        s_dict = {
            "student_id": sid,
            "student_name": name,
            "student_level": norm_level,
            "raw_level": str(raw_level).strip() if raw_level else norm_level,
            "batch_type": btype,
            "batch_name": batch_name,
            "group_number": grp_no,
            "required_classes": planned_cnt,
            "planned_classes": planned_cnt,
            "actual_classes": len(scheduled_dates),
            "region_timezone": "IST",
            "mkca_rating": rating_flt,
            "assigned_batch_id": batch_id,
            "mon_pref": "Available" if "Monday" in p_days else "Not Available",
            "tue_pref": "Available" if "Tuesday" in p_days else "Not Available",
            "wed_pref": "Available" if "Wednesday" in p_days else "Not Available",
            "thu_pref": "Available" if "Thursday" in p_days else "Not Available",
            "fri_pref": "Available" if "Friday" in p_days else "Not Available",
            "sat_pref": "Available" if "Saturday" in p_days else "Not Available",
            "sun_pref": "Available" if "Sunday" in p_days else "Not Available",
            "additional_comments": combined_notes,
            "recurring_slots": recurring_slots
        }
        students.append(s_dict)
        student_map[sid] = s_dict

    # 2. Build ScheduledClass list from sessions
    scheduled_classes = []
    # Sort sessions strictly chronologically
    sorted_sess_keys = sorted(
        sessions.keys(),
        key=lambda k: (parse_time_slot_sort_key(k[0], k[1]), k[2])
    )

    for sess_key in sorted_sess_keys:
        s_info = sessions[sess_key]
        date_str = s_info["date"]
        time_slot = s_info["time_slot"]
        coach_name = s_info["coach_name"]
        
        # Primary batch and level
        distinct_batches = sorted(list(s_info["batches"]))
        distinct_levels = sorted(list(s_info["levels"]))
        distinct_btypes = s_info["batch_types"]

        batch_name = " / ".join(distinct_batches) if len(distinct_batches) <= 2 else distinct_batches[0]
        level_name = distinct_levels[0] if len(distinct_levels) == 1 else "MIXED"
        batch_type = "G" if "G" in distinct_btypes else ("L" if "L" in distinct_btypes else "I")

        clean_date = date_str.replace("-", "")
        clean_time = re.sub(r'[^a-zA-Z0-9]', '', time_slot[:8])
        clean_coach = re.sub(r'[^a-zA-Z0-9]', '', coach_name)
        class_id = f"CLS_{clean_date}_{clean_time}_{clean_coach}"

        # Ensure unique class_id in list
        counter = 1
        base_id = class_id
        existing_ids = {c["class_id"] for c in scheduled_classes}
        while class_id in existing_ids:
            class_id = f"{base_id}_{counter}"
            counter += 1

        cls_item = {
            "class_id": class_id,
            "date": date_str,
            "day": s_info["day"],
            "time_slot": time_slot,
            "coach_name": coach_name,
            "batch_name": batch_name,
            "student_level": level_name,
            "batch_type": batch_type,
            "student_ids": s_info["student_ids"],
            "student_names": s_info["student_names"],
            "warnings": [],
            "is_manual_override": False
        }
        scheduled_classes.append(cls_item)

    # 3. Build Coaches List
    coaches_list = []
    for c_name in sorted(list(coaches_set)):
        coaches_list.append({
            "coach_name": c_name,
            "levels_handled": [
                "Basic 1", "Basic 2", "Beginner 1", "Beginner 2",
                "Early Intermediate 1", "Early Intermediate 2",
                "Intermediate 1", "Intermediate 2"
            ],
            "monthly_capacity_min": 0,
            "monthly_capacity_max": 200,
            "mon_max": 10, "tue_max": 10, "wed_max": 10, "thu_max": 10, "fri_max": 10,
            "sat_max": 15, "sun_max": 10,
            "sunday_pref": "Available",
            "preferred_timings": "All Operating Hours"
        })

    # 4. Build Coach Communication Slots (Output 1)
    coach_sched_map = defaultdict(list)
    for c in scheduled_classes:
        k = (c["date"], c["day"], c["time_slot"])
        if c["coach_name"] not in coach_sched_map[k]:
            coach_sched_map[k].append(c["coach_name"])

    coach_schedule_slots = [
        {
            "date": k[0],
            "day": k[1],
            "time_slot": k[2],
            "coaches": v
        }
        for k, v in sorted(coach_sched_map.items(), key=lambda item: parse_time_slot_sort_key(item[0][0], item[0][2]))
    ]

    # 5. Open SQLite Database and Save Master Data First
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    # Clear old data
    cur.execute("DELETE FROM students")
    cur.execute("DELETE FROM coaches")
    cur.execute("DELETE FROM batches")
    cur.execute("DELETE FROM schedules")
    cur.execute("DELETE FROM active_metadata")

    # Insert Students
    for s in students:
        cur.execute("""
            INSERT INTO students (student_id, student_name, student_level, batch_type, region_timezone, required_classes, data_json)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (
            s["student_id"],
            s["student_name"],
            s["student_level"],
            s["batch_type"],
            s["region_timezone"],
            s["required_classes"],
            json.dumps(s)
        ))

    # Insert Coaches
    for c in coaches_list:
        cur.execute("""
            INSERT INTO coaches (coach_name, levels_handled_json, monthly_capacity_min, monthly_capacity_max, data_json)
            VALUES (?, ?, ?, ?, ?)
        """, (
            c["coach_name"],
            json.dumps(c["levels_handled"]),
            c["monthly_capacity_min"],
            c["monthly_capacity_max"],
            json.dumps(c)
        ))

    # Insert Batches
    for b_id, b in batches_map.items():
        b["student_count"] = len(b["students"])
        cur.execute("""
            INSERT INTO batches (batch_id, batch_name, level, batch_type, fixed_trainer, schedule_timings, student_ids_json, data_json)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            b["batch_id"],
            b["batch_name"],
            b["level"],
            b["batch_type"],
            b["fixed_trainer"],
            b["schedule_timings"],
            json.dumps(b["student_ids"]),
            json.dumps(b)
        ))

    conn.commit()

    # 6. Compute Master Data Fingerprint now that master tables are populated
    current_fp = compute_master_data_fingerprint(db_path=DB_PATH)
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %I:%M %p")

    # Build Complete ScheduleResult Dict
    schedule_result_dict = {
        "schedule_id": "SCH_OCT2026_FINAL",
        "status": "Finalized",
        "start_date": "2026-10-01",
        "end_date": "2026-10-31",
        "total_students_considered": len(students),
        "successfully_scheduled_students": len(students),
        "unscheduled_students_count": 0,
        "accountability_passed": True,
        "scheduled_classes": scheduled_classes,
        "unscheduled_records": [],
        "coach_schedule": coach_schedule_slots,
        "master_data_fingerprint": current_fp,
        "is_stale": False,
        "created_at": now_iso
    }

    # Insert Active Schedule
    cur.execute("""
        INSERT INTO schedules (schedule_id, status, start_date, end_date, total_students, scheduled_students, unscheduled_students, accountability_passed, data_json, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        schedule_result_dict["schedule_id"],
        schedule_result_dict["status"],
        schedule_result_dict["start_date"],
        schedule_result_dict["end_date"],
        schedule_result_dict["total_students_considered"],
        schedule_result_dict["successfully_scheduled_students"],
        schedule_result_dict["unscheduled_students_count"],
        1,
        json.dumps(schedule_result_dict),
        schedule_result_dict["created_at"]
    ))

    # Insert Active Metadata
    cur.execute("INSERT OR REPLACE INTO active_metadata (key, value) VALUES ('latest_schedule_id', 'SCH_OCT2026_FINAL')")
    cur.execute("INSERT OR REPLACE INTO active_metadata (key, value) VALUES ('active_schedule_id', 'SCH_OCT2026_FINAL')")
    cur.execute("INSERT OR REPLACE INTO active_metadata (key, value) VALUES ('master_cleared', 'false')")
    cur.execute("INSERT OR REPLACE INTO active_metadata (key, value) VALUES ('last_filename', 'Oct26_Schedule_Final.xlsx')")
    cur.execute("INSERT OR REPLACE INTO active_metadata (key, value) VALUES ('last_upload_timestamp', '2026-10-01 12:00 AM')")

    conn.commit()
    cur.execute("VACUUM")
    conn.close()

    # 7. Write pre-compiled reference JSON cache
    oct_exact_by_date = defaultdict(list)
    for cls in scheduled_classes:
        oct_exact_by_date[cls["date"]].append(cls)

    # Weekday templates
    weekday_templates = defaultdict(list)
    for s in students:
        for d_name, (ts, cn) in s["recurring_slots"].items():
            weekday_templates[d_name].append({
                "student_id": s["student_id"],
                "student_name": s["student_name"],
                "time_slot": ts,
                "coach_name": cn,
                "batch_name": s["batch_name"],
                "student_level": s["student_level"],
                "batch_type": s["batch_type"]
            })

    json_payload = {
        "source": "Oct26_Schedule_Final.xlsx",
        "students_raw": students,
        "coaches_list": sorted(list(coaches_set)),
        "oct_date_range": ["2026-10-01", "2026-10-31"],
        "oct_exact_classes_by_date": oct_exact_by_date,
        "weekday_templates": weekday_templates
    }
    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(json_payload, f, indent=2)

    print("\n========================================================")
    print("SUCCESSFULLY IMPORTED AUTHORITATIVE OCTOBER 2026 DATA:")
    print(f"  - Total Enrolled Students : {len(students)}")
    print(f"  - Total Scheduled Classes : {len(scheduled_classes)}")
    print(f"  - Total Active Batches    : {len(batches_map)}")
    print(f"  - Total Coaches Active    : {len(coaches_set)}")
    print(f"  - Active Schedule ID      : SCH_OCT2026_FINAL (Status: Finalized)")
    print(f"  - Master Fingerprint      : {current_fp[:16]}...")
    print(f"  - Database Updated        : {DB_PATH}")
    print(f"  - Pre-compiled JSON Cache : {JSON_PATH}")
    print("========================================================")

if __name__ == "__main__":
    parse_and_feed()
