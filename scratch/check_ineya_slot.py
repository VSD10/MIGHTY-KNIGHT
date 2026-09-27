import openpyxl

wb = openpyxl.load_workbook(r"sample_data/Oct'26 Schedule_FRESH-1.xlsx", data_only=True)
ws = wb.active
for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 1).value
    batch = ws.cell(r, 5).value
    level = ws.cell(r, 4).value
    for c in range(8, ws.max_column, 2):
        d = ws.cell(3, c).value
        if d and hasattr(d, 'strftime') and d.strftime('%Y-%m-%d') == '2026-10-03':
            t = str(ws.cell(r, c).value or '')
            tr = str(ws.cell(r, c+1).value or '')
            if 'prakash' in tr.lower() and ('19' in t or '7' in t):
                print(f"Same slot: {name} | Level={level} | Batch={batch} | Time={t}")
