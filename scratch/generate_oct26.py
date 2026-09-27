"""
October 2026 Chess Schedule Excel Generator
============================================
Reads the Oct'26 Schedule_FRESH-1.xlsx template (which already has all student 
time/trainer cells populated) and:
  1. Counts each student's actual assignments (non-empty time cells)
  2. Fills in the per-student "Actual Classes" column (col 7)
  3. Recalculates daily group-class counts (unique Date+Time+Trainer)
  4. Verifies row 2 group-class totals match the computed values
  5. Copies styles faithfully and saves as Oct'26_Schedule_Generated.xlsx

Layout (confirmed from inspection):
  Row 1  : day-name headers (cols 8,10,12,...68) + "Comments" (col 70)
  Row 2  : "Actual Classes" label (col 7) + daily group-class totals (cols 8,10,...68)
  Row 3  : column headers (Student Name, Stud ID, MKCA Rating, Level, Batch, date cells)
  Rows 4+ : student data
             col 1  = Student Name
             col 2  = Stud ID
             col 3  = MKCA Rating
             col 4  = Level
             col 5  = Batch
             col 6  = Planned Classes
             col 7  = Actual Classes (to be filled)
             col 8,9   = time/trainer for Oct 1
             col 10,11 = time/trainer for Oct 2
             ...
             col 68,69 = time/trainer for Oct 31
             col 70 = Comments
"""

import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
import os
import datetime
import copy

# ── paths ─────────────────────────────────────────────────────────────────────
BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
TEMPLATE_PATH = os.path.join(BASE_DIR, 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
OUTPUT_PATH   = os.path.join(BASE_DIR, 'sample_data', "Oct'26_Schedule_Generated.xlsx")

# ── load workbook ─────────────────────────────────────────────────────────────
print(f"Loading template: {TEMPLATE_PATH}")
wb = load_workbook(TEMPLATE_PATH)
ws = wb.active
print(f"Sheet: {ws.title}  |  Rows: {ws.max_row}  |  Cols: {ws.max_column}")

# ── build date→column index map ───────────────────────────────────────────────
# Even cols 8..68 = date; odd cols 9..69 = trainer (paired)
date_cols = {}   # date_obj -> (time_col, trainer_col)
for col in range(8, 70, 2):
    cell_val = ws.cell(3, col).value
    if isinstance(cell_val, datetime.datetime):
        date_cols[cell_val.date()] = (col, col + 1)

print(f"Dates mapped: {sorted(date_cols.keys())}")

# ── STEP 1: count per-student actual classes & fill col 7 ─────────────────────
print("\nFilling per-student Actual Classes (col 7)...")

student_actuals = {}  # row -> count

for row in range(4, ws.max_row + 1):
    name = ws.cell(row, 1).value
    if not name:
        continue
    
    count = 0
    for col in range(8, 70, 2):     # time columns
        time_val = ws.cell(row, col).value
        if time_val is not None:
            count += 1
    
    student_actuals[row] = count
    ws.cell(row, 7).value = count

print(f"  Filled {len(student_actuals)} students")

# Quick sanity sample
for row in list(student_actuals.keys())[:5]:
    name = ws.cell(row, 1).value
    planned = ws.cell(row, 6).value
    actual  = student_actuals[row]
    status  = "OK" if (planned is None or actual <= planned) else "!! EXCEEDS"
    print(f"  Row {row}: {name:30s}  planned={planned}  actual={actual}  {status}")

# ── STEP 2: compute daily group-class counts ───────────────────────────────────
print("\nComputing daily group-class counts...")

date_sorted = sorted(date_cols.keys())
daily_groups = {}  # date -> count

for d in date_sorted:
    time_col, trainer_col = date_cols[d]
    groups = set()
    for row in range(4, ws.max_row + 1):
        name = ws.cell(row, 1).value
        if not name:
            continue
        time_val    = ws.cell(row, time_col).value
        trainer_val = ws.cell(row, trainer_col).value
        if time_val is not None and trainer_val is not None:
            groups.add((time_val, trainer_val))
    daily_groups[d] = len(groups)

print("  Date          Computed  Template  Match")
print("  " + "-" * 45)
all_match = True
for d in date_sorted:
    time_col, _ = date_cols[d]
    template_count = ws.cell(2, time_col).value
    computed_count = daily_groups[d]
    match = "OK" if template_count == computed_count else f"!! DIFF (template={template_count})"
    if template_count != computed_count:
        all_match = False
    print(f"  {d}     {computed_count:>4}      {template_count!s:>5}   {match}")

if all_match:
    print("\n  [OK] All daily group-class counts match template!")
else:
    print("\n  [WARN] Some counts differ -- template values kept (template is authoritative)")

# ── STEP 3: verify per-student caps ───────────────────────────────────────────
print("\nValidating Planned Classes caps...")
violations = []
for row, actual in student_actuals.items():
    name    = ws.cell(row, 1).value
    planned = ws.cell(row, 6).value
    if planned is not None and actual > planned:
        violations.append((name, planned, actual))

if violations:
    print(f"  [WARN] {len(violations)} violations:")
    for name, planned, actual in violations:
        print(f"    {name}: planned={planned}, actual={actual}")
else:
    print("  [OK] No cap violations")

# ── STEP 4: apply basic formatting to col 7 ──────────────────────────────────
print("\nApplying formatting to Actual Classes column...")

# Copy style from col 6 (Planned Classes) to col 7 for student rows
for row in range(4, ws.max_row + 1):
    if ws.cell(row, 1).value:
        src_cell = ws.cell(row, 6)
        dst_cell = ws.cell(row, 7)
        # Copy font and alignment
        if src_cell.font:
            dst_cell.font = copy.copy(src_cell.font)
        if src_cell.alignment:
            dst_cell.alignment = copy.copy(src_cell.alignment)
        if src_cell.fill:
            dst_cell.fill = copy.copy(src_cell.fill)
        if src_cell.border:
            dst_cell.border = copy.copy(src_cell.border)
        dst_cell.number_format = '0'

print("  Done")

# ── STEP 5: summary ───────────────────────────────────────────────────────────
total_assignments = sum(student_actuals.values())
total_group_classes = sum(daily_groups.values())
print(f"\n{'='*55}")
print(f"Summary for October 2026")
print(f"{'='*55}")
print(f"  Students:                  {len(student_actuals)}")
print(f"  Total student assignments: {total_assignments}")
print(f"  Total group classes:       {total_group_classes}")
print(f"  Date range:                {date_sorted[0]} – {date_sorted[-1]}")
print(f"{'='*55}")

# ── STEP 6: save ──────────────────────────────────────────────────────────────
print(f"\nSaving to: {OUTPUT_PATH}")
wb.save(OUTPUT_PATH)
print("[OK] Saved successfully!")
print(f"\nOutput file: {OUTPUT_PATH}")
