"""
Quick verify: show all student actuals and check a few row samples
"""
import openpyxl, os, datetime

path = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26_Schedule_Generated.xlsx"))
wb = openpyxl.load_workbook(path)
ws = wb.active

print("=== VERIFICATION: Oct'26_Schedule_Generated.xlsx ===\n")

# Row 2 - Actual Classes label and daily totals
print("Row 2 (group class totals row):")
print(f"  Col 7 label: {ws.cell(2,7).value}")
for col in range(8, 70, 2):
    d = ws.cell(3, col).value
    g = ws.cell(2, col).value
    if d:
        print(f"  {d.strftime('%m/%d')} -> {g} group classes")

print()
print("Per-student Actual Classes (col 7):")
violations = []
for row in range(4, ws.max_row + 1):
    name = ws.cell(row, 1).value
    if not name:
        continue
    planned = ws.cell(row, 6).value
    actual  = ws.cell(row, 7).value
    status = "OK" if (planned is None or actual <= planned) else "!! OVER LIMIT"
    if planned is not None and actual > planned:
        violations.append((name, planned, actual))
    print(f"  {name:40s} planned={str(planned):>4}  actual={str(actual):>3}  {status}")

print()
print(f"Total students: {ws.max_row - 3}")
print(f"Violations: {len(violations)}")
if violations:
    for v in violations:
        print(f"  !! {v[0]}: planned={v[1]}, actual={v[2]}")
else:
    print("  NONE - all within Planned Classes limits")
