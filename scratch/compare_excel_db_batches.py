import sqlite3
import json
import openpyxl

conn = sqlite3.connect('backend/data/chess_scheduler.db')
cursor = conn.cursor()

# Get student_id -> batch_name from batches table
batch_mapping = {}
for b in cursor.execute("SELECT batch_id, batch_name, level, student_ids_json FROM batches").fetchall():
    b_id, b_name, b_level, sids_json = b
    sids = json.loads(sids_json) if sids_json else []
    for sid in sids:
        batch_mapping[sid] = b_name

wb = openpyxl.load_workbook(r"sample_data/Oct'26 Schedule_FRESH-1.xlsx", data_only=True)
ws = wb.active

mismatches = []
matched = 0

for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if not name or not str(name).strip():
        continue
    sid = str(ws.cell(r, 2).value or f"MKS{r:05d}").strip()
    excel_batch = str(ws.cell(r, 5).value or "").strip()
    
    db_batch = batch_mapping.get(sid)
    if db_batch != excel_batch:
        mismatches.append((sid, name, excel_batch, db_batch))
    else:
        matched += 1

print(f"Excel vs DB Batches: {matched} matched, {len(mismatches)} mismatches")
if mismatches:
    print("Mismatches sample:", mismatches[:10])

conn.close()
