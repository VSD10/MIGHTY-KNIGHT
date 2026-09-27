import openpyxl

wb = openpyxl.load_workbook(r"sample_data/Oct'26 Schedule_FRESH-1.xlsx", data_only=True)
ws = wb.active
for r in range(4, ws.max_row + 1):
    name = str(ws.cell(r, 1).value or '')
    if 'ineya' in name.lower():
        sid = ws.cell(r, 2).value
        rating = ws.cell(r, 3).value
        level = ws.cell(r, 4).value
        batch = ws.cell(r, 5).value
        classes = ws.cell(r, 6).value
        print(f"Excel Row {r}: Name='{name}', ID='{sid}', Rating='{rating}', Level='{level}', Batch='{batch}', Classes='{classes}'")
        for c in range(8, ws.max_column, 2):
            t = ws.cell(r, c).value
            tr = ws.cell(r, c+1).value
            d = ws.cell(3, c).value
            if t and tr:
                d_str = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)
                print(f"   {d_str}: {t} with {tr}")
