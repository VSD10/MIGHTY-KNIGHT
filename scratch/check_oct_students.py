import openpyxl
import os

path = os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
wb = openpyxl.load_workbook(os.path.abspath(path), data_only=True)
ws = wb.active

oct_students = []
for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if name:
        oct_students.append((str(name).strip(), ws.cell(r, 2).value))

print(f"Found {len(oct_students)} students in Oct'26 Schedule_FRESH-1.xlsx")
print("First 5:", oct_students[:5])
print("Last 5:", oct_students[-5:])
