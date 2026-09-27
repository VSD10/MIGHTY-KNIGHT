import openpyxl
import os

path = os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
wb = openpyxl.load_workbook(os.path.abspath(path), data_only=True)
ws = wb.active

row2_total = 0
for col in range(8, ws.max_column, 2):
    v = ws.cell(2, col).value
    if v is not None and isinstance(v, (int, float)):
        row2_total += int(v)

print(f"Sum of Row 2 Actual Classes in Excel: {row2_total}")
