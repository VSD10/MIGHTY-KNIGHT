import os
import openpyxl
from collections import defaultdict

fpath = os.path.join("sample_data", "Oct'26 Schedule_FRESH-1.xlsx")
wb = openpyxl.load_workbook(fpath, data_only=True)
print("Sheet names:", wb.sheetnames)
ws = wb.active
print("Active sheet:", ws.title)

day_coaches = defaultdict(lambda: defaultdict(int))
for r in range(4, 131):
    for c in range(8, ws.max_column, 2):
        d_val = ws.cell(3, c).value
        if not d_val: continue
        d_str = d_val.strftime('%Y-%m-%d')
        t_val = ws.cell(r, c).value
        tr_val = ws.cell(r, c+1).value
        if t_val and tr_val:
            day_coaches[d_str][str(tr_val).strip()] += 1

print("\n--- Coaches per Day in Excel ---")
for d in sorted(day_coaches.keys())[:10]:
    print(d, dict(day_coaches[d]))
