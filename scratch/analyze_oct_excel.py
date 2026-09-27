import openpyxl
import os
import datetime
from collections import defaultdict, Counter

path = os.path.join(os.path.dirname(__file__), '..', 'sample_data', "Oct'26 Schedule_FRESH-1.xlsx")
wb = openpyxl.load_workbook(os.path.abspath(path), data_only=True)
ws = wb.active

date_map = {}
for col in range(8, ws.max_column, 2):
    d_val = ws.cell(3, col).value
    day_val = ws.cell(1, col).value
    if d_val:
        if isinstance(d_val, datetime.datetime):
            d_obj = d_val.date()
        else:
            d_obj = d_val
        date_map[col] = (d_obj, day_val, col + 1)

students = []
student_weekday_patterns = defaultdict(lambda: defaultdict(set))
trainers_found = set()

for r in range(4, ws.max_row + 1):
    name = ws.cell(r, 1).value
    if not name or not str(name).strip():
        continue
    
    stud_id = ws.cell(r, 2).value
    rating = ws.cell(r, 3).value
    level = ws.cell(r, 4).value
    batch = ws.cell(r, 5).value
    planned = ws.cell(r, 6).value
    comments = ws.cell(r, ws.max_column).value
    
    assignments = []
    weekday_slots = defaultdict(list)
    
    for col, (d_obj, day_val, tr_col) in date_map.items():
        t_val = ws.cell(r, col).value
        tr_val = ws.cell(r, tr_col).value
        if t_val is not None:
            # format time
            if isinstance(t_val, datetime.time):
                time_str = t_val.strftime("%I:%M %p").lstrip('0')
            else:
                time_str = str(t_val)
            
            trainer_str = str(tr_val).strip() if tr_val else "Unassigned"
            trainers_found.add(trainer_str)
            assignments.append((d_obj, day_val, time_str, trainer_str))
            
            # Record weekday recurrence
            # day_val can be Thursday, Friday, etc.
            weekday_slots[day_val].append((time_str, trainer_str))
            
    students.append({
        'row': r,
        'name': str(name).strip(),
        'stud_id': str(stud_id).strip() if stud_id else f"STU_{r}",
        'rating': rating,
        'level': str(level).strip() if level else "Basic",
        'batch': str(batch).strip() if batch else "General",
        'planned': planned,
        'comments': comments,
        'assignments': assignments,
        'weekday_slots': dict(weekday_slots)
    })

print(f"Total students: {len(students)}")
total_assignments = sum(len(s['assignments']) for s in students)
print(f"Total student class assignments: {total_assignments}")
print(f"Trainers found ({len(trainers_found)}): {sorted(trainers_found)}")

# Group classes count (unique date + time + trainer)
unique_classes_per_date = defaultdict(set)
for s in students:
    for (d_obj, day_val, time_str, trainer_str) in s['assignments']:
        unique_classes_per_date[d_obj].add((time_str, trainer_str))

print("Total unique classes across October:", sum(len(v) for v in unique_classes_per_date.values()))
print("\nSample student recurring patterns:")
for s in students[:5]:
    print(f"Student: {s['name']} ({s['stud_id']}) | Planned: {s['planned']} | Oct Classes: {len(s['assignments'])}")
    for day, slots in s['weekday_slots'].items():
        slot_counts = Counter(slots)
        patterns_str = ", ".join([f"{t} with {tr} (x{c})" for (t, tr), c in slot_counts.items()])
        print(f"   {day}: {patterns_str}")
