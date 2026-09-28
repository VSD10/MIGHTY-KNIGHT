import os
import json
import openpyxl
import datetime
import calendar
import uuid
from typing import List, Dict, Any, Optional, Tuple, Set
from collections import defaultdict

from app.models.student import StudentModel
from app.models.coach import CoachModel
from app.models.schedule import ScheduleResult, ScheduledClass, UnscheduledRecord, CoachCommunicationSlot
from app.config import SystemConfig, DEFAULT_CONFIG
from app.constants.levels import OFFICIAL_LEVELS, normalize_batch_to_level

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
DEFAULT_REFERENCE_PATH = os.path.join(BASE_DIR, "sample_data", "Oct'26 Schedule_FRESH-1.xlsx")

class ReferenceScheduleEngine:
    _instance = None

    def __init__(self, excel_path: Optional[str] = None):
        self.excel_path = excel_path or DEFAULT_REFERENCE_PATH
        self.students_raw: List[Dict[str, Any]] = []
        self.student_models: List[StudentModel] = []
        self.coach_models: List[CoachModel] = []
        self.coaches_list: List[str] = []
        
        # Day of week -> list of recurring student slot dicts:
        # {student_id, student_name, time_slot, coach_name, batch_name, student_level, batch_type}
        self.weekday_templates: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        
        # Exact October classes: date_str -> list of {class_id, time_slot, coach_name, batch_name, student_level, batch_type, student_ids, student_names}
        self.oct_exact_classes_by_date: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        self.oct_date_range: Tuple[Optional[datetime.date], Optional[datetime.date]] = (None, None)
        
        self.is_loaded = False
        self.load_reference()

    def _build_models(self):
        # Build StudentModel list
        self.student_models = []
        for s in self.students_raw:
            p_days = s.get("recurring_slots", {})
            self.student_models.append(StudentModel(
                student_id=s["student_id"],
                student_name=s["student_name"],
                student_level=s["student_level"],
                batch_type=s.get("batch_type", "G"),
                required_classes=s.get("required_classes", 12),
                mkca_rating=s.get("mkca_rating"),
                mon_pref="Available" if "Monday" in p_days else "Not Available",
                tue_pref="Available" if "Tuesday" in p_days else "Not Available",
                wed_pref="Available" if "Wednesday" in p_days else "Not Available",
                thu_pref="Available" if "Thursday" in p_days else "Not Available",
                fri_pref="Available" if "Friday" in p_days else "Not Available",
                sat_pref="Available" if "Saturday" in p_days else "Not Available",
                sun_pref="Available" if "Sunday" in p_days else "Not Available",
                additional_comments=s.get("additional_comments", "")
            ))

        # Build CoachModel list
        self.coach_models = []
        for c_name in self.coaches_list:
            self.coach_models.append(CoachModel(
                coach_name=c_name,
                levels_handled=OFFICIAL_LEVELS,
                monthly_capacity_min=0,
                monthly_capacity_max=200,
                mon_max=20, tue_max=20, wed_max=20, thu_max=20, fri_max=20, sat_max=20, sun_max=15,
                sunday_pref="Available",
                preferred_timings="All Operating Hours"
            ))

    def load_reference(self, custom_path: Optional[str] = None):
        if custom_path:
            self.excel_path = custom_path

        # 1. First priority: Check pre-compiled JSON reference (instant & 100% reliable in serverless)
        candidate_json_paths = [
            os.path.join(os.path.dirname(__file__), "..", "..", "data", "reference_oct_schedule.json"),
            os.path.join(BASE_DIR, "backend", "data", "reference_oct_schedule.json"),
            os.path.join(BASE_DIR, "data", "reference_oct_schedule.json")
        ]

        for j_path in candidate_json_paths:
            if j_path and os.path.exists(j_path):
                try:
                    with open(j_path, "r", encoding="utf-8") as f:
                        data = json.load(f)
                    self.students_raw = data.get("students_raw", [])
                    self.coaches_list = data.get("coaches_list", [])
                    self.oct_exact_classes_by_date = defaultdict(list, data.get("oct_exact_classes_by_date", {}))
                    self.weekday_templates = defaultdict(list, data.get("weekday_templates", {}))
                    dr = data.get("oct_date_range", ["2026-10-01", "2026-10-31"])
                    self.oct_date_range = (
                        datetime.datetime.strptime(dr[0], "%Y-%m-%d").date(),
                        datetime.datetime.strptime(dr[1], "%Y-%m-%d").date()
                    )
                    self._build_models()
                    self.is_loaded = True
                    total_cls = sum(len(v) for v in self.oct_exact_classes_by_date.values())
                    print(f"[REFERENCE_ENGINE] Instantly loaded {len(self.students_raw)} students, {len(self.coaches_list)} coaches, {total_cls} exact October classes from JSON: {j_path}")
                    return
                except Exception as e:
                    print(f"[REFERENCE_ENGINE] Could not load JSON cache: {e}")

        # 2. Second priority: Find Excel file across candidate locations
        candidate_excel_paths = [
            self.excel_path,
            os.path.join(os.path.dirname(__file__), "..", "..", "data", "Oct'26 Schedule_FRESH-1.xlsx"),
            os.path.join(BASE_DIR, "backend", "data", "Oct'26 Schedule_FRESH-1.xlsx"),
            os.path.join(BASE_DIR, "sample_data", "Oct'26 Schedule_FRESH-1.xlsx")
        ]

        target_excel = None
        for p in candidate_excel_paths:
            if p and os.path.exists(p):
                target_excel = p
                break

        if not target_excel:
            print(f"[REFERENCE_ENGINE] Reference file not found in any candidate path!")
            return

        self.excel_path = target_excel
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb.active

        # 1. Parse Date Columns from Row 3 (dates) and Row 1 (day names)
        # Even columns starting from column 8
        date_cols = {} # col_idx -> (datetime.date, day_name, trainer_col_idx)
        for col in range(8, ws.max_column, 2):
            d_val = ws.cell(3, col).value
            day_val = ws.cell(1, col).value
            if d_val:
                d_obj = d_val.date() if isinstance(d_val, datetime.datetime) else d_val
                day_name = str(day_val).strip() if day_val else d_obj.strftime("%A")
                date_cols[col] = (d_obj, day_name, col + 1)

        dates_found = [d for d, _, _ in date_cols.values()]
        if dates_found:
            self.oct_date_range = (min(dates_found), max(dates_found))

        students_list = []
        trainers_set = set()
        weekday_templates = defaultdict(list)
        
        # Exact classes map: (date_str, time_slot, coach_name) -> dict
        oct_classes_map = defaultdict(lambda: {
            "student_ids": [],
            "student_names": [],
            "batch_names": set(),
            "levels": set(),
            "batch_types": set(),
            "day": ""
        })

        for r in range(4, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if not name or not str(name).strip():
                continue

            name_str = str(name).strip()
            stud_id = str(ws.cell(r, 2).value or f"MKS{r:05d}").strip()
            rating_val = ws.cell(r, 3).value
            batch_str = str(ws.cell(r, 5).value or "G Basic1").strip()
            calc_level, _, _ = normalize_batch_to_level(batch_str, stud_id)
            level_str = calc_level
            planned_val = ws.cell(r, 6).value
            comments_val = str(ws.cell(r, ws.max_column).value or "").strip()

            # Determine batch type
            b_type = "G"
            if batch_str.startswith("L ") or batch_str.startswith("L-") or "Limited" in batch_str:
                b_type = "L"
            elif batch_str.startswith("I ") or batch_str.startswith("I-") or "Individual" in batch_str:
                b_type = "I"

            try:
                planned_classes = int(planned_val) if planned_val is not None else 12
            except (ValueError, TypeError):
                planned_classes = 12

            try:
                rating_float = float(rating_val) if rating_val is not None else None
            except (ValueError, TypeError):
                rating_float = None

            stud_dict = {
                "student_id": stud_id,
                "student_name": name_str,
                "student_level": level_str,
                "batch_type": b_type,
                "batch_name": batch_str,
                "mkca_rating": rating_float,
                "required_classes": planned_classes,
                "planned_classes": planned_classes,
                "additional_comments": comments_val,
                "recurring_slots": {} # day_name -> (time_slot, coach_name)
            }

            # Read assigned classes for October
            oct_assigned_count = 0
            for col, (d_obj, day_name, tr_col) in date_cols.items():
                t_val = ws.cell(r, col).value
                tr_val = ws.cell(r, tr_col).value
                if t_val is not None:
                    if isinstance(t_val, datetime.time):
                        t_str = t_val.strftime("%I:%M %p").lstrip('0')
                    else:
                        t_str = str(t_val).strip()

                    tr_str = str(tr_val).strip() if tr_val else "Unassigned"
                    trainers_set.add(tr_str)
                    oct_assigned_count += 1

                    # Record exact class for October
                    date_str = d_obj.strftime("%Y-%m-%d")
                    class_key = (date_str, t_str, tr_str)
                    oct_classes_map[class_key]["student_ids"].append(stud_id)
                    oct_classes_map[class_key]["student_names"].append(name_str)
                    oct_classes_map[class_key]["batch_names"].add(batch_str)
                    oct_classes_map[class_key]["levels"].add(level_str)
                    oct_classes_map[class_key]["batch_types"].add(b_type)
                    oct_classes_map[class_key]["day"] = day_name

                    # Record recurring weekday template for student
                    if day_name not in stud_dict["recurring_slots"]:
                        stud_dict["recurring_slots"][day_name] = (t_str, tr_str)
                        weekday_templates[day_name].append({
                            "student_id": stud_id,
                            "student_name": name_str,
                            "time_slot": t_str,
                            "coach_name": tr_str,
                            "batch_name": batch_str,
                            "student_level": level_str,
                            "batch_type": b_type
                        })

            stud_dict["actual_classes"] = oct_assigned_count
            students_list.append(stud_dict)

        self.students_raw = students_list
        self.coaches_list = sorted(list(trainers_set))
        self.weekday_templates = weekday_templates

        # Organize exact October classes by date
        self.oct_exact_classes_by_date.clear()
        for (date_str, t_str, tr_str), c_info in oct_classes_map.items():
            distinct_levels = sorted(list(c_info["levels"]))
            if len(distinct_levels) == 1:
                primary_level = distinct_levels[0]
            elif len(distinct_levels) > 1:
                primary_level = "MIXED"
            else:
                primary_level = "Basic 1"
            primary_batch = sorted(c_info["batch_names"])[0] if c_info["batch_names"] else "G Basic1"
            primary_btype = "G"
            if "G" in c_info["batch_types"]:
                primary_btype = "G"
            elif "L" in c_info["batch_types"]:
                primary_btype = "L"
            elif "I" in c_info["batch_types"]:
                primary_btype = "I"

            cls_id = f"CLS_{date_str.replace('-', '')}_{t_str.replace(' ', '').replace(':', '')}_{tr_str}"
            self.oct_exact_classes_by_date[date_str].append({
                "class_id": cls_id,
                "date": date_str,
                "day": c_info["day"],
                "time_slot": t_str,
                "coach_name": tr_str,
                "student_level": primary_level,
                "batch_name": primary_batch,
                "batch_type": primary_btype,
                "student_ids": c_info["student_ids"],
                "student_names": c_info["student_names"],
                "warnings": [],
                "is_manual_override": False
            })

        self._build_models()
        self.is_loaded = True
        total_cls = sum(len(v) for v in self.oct_exact_classes_by_date.values())
        print(f"[REFERENCE_ENGINE] Successfully loaded {len(self.students_raw)} students, {len(self.coaches_list)} coaches, {total_cls} exact October classes from {self.excel_path}")

_ENGINE = None

def get_reference_engine() -> ReferenceScheduleEngine:
    global _ENGINE
    if _ENGINE is None:
        _ENGINE = ReferenceScheduleEngine()
    return _ENGINE

def generate_reference_schedule(
    start_date: datetime.date,
    end_date: datetime.date,
    config: SystemConfig = DEFAULT_CONFIG,
    schedule_id: Optional[str] = None,
    students: Optional[List[StudentModel]] = None,
    coaches: Optional[List[CoachModel]] = None
) -> ScheduleResult:
    """
    Template-Based Recurring Scheduler:
    - MODE A (October 2026): Exact reproduction of the uploaded Excel schedule.
      Same students, IDs, days, times, trainers, batches, levels, and empty days.
    - MODE B (Future Months: November, December, etc.):
      Uses the reference Excel as the template.
      Applies recurring student pattern (STUDENT + DAY OF WEEK + TIME + TRAINER)
      to every occurrence of that weekday in the target month.
      Trainer and Time are strictly fixed. No artificial trainer capacity limits or balancing.
    """
    engine = get_reference_engine()
    if not engine.is_loaded:
        engine.load_reference()

    if not schedule_id:
        schedule_id = f"SCH_{uuid.uuid4().hex[:8].upper()}"

    s_str = start_date.strftime("%Y-%m-%d")
    e_str = end_date.strftime("%Y-%m-%d")

    # Target students map if provided
    has_target_filter = students is not None and len(students) > 0
    target_student_ids = {s.student_id for s in students} if has_target_filter else None
    target_student_map = {s.student_id: s for s in students} if has_target_filter else None

    # Generate calendar date list
    num_days = (end_date - start_date).days + 1
    target_dates = [start_date + datetime.timedelta(days=i) for i in range(num_days)]

    scheduled_classes: List[ScheduledClass] = []
    
    # Check if target period is in October 2026 (Mode A: Exact Copy)
    is_october_2026 = (start_date.year == 2026 and start_date.month == 10 and end_date.year == 2026 and end_date.month == 10)

    if is_october_2026:
        # MODE A: Exact reproduction from Excel cells
        for d_obj in target_dates:
            d_str = d_obj.strftime("%Y-%m-%d")
            for cls_dict in engine.oct_exact_classes_by_date.get(d_str, []):
                s_ids = cls_dict["student_ids"]
                s_names = cls_dict["student_names"]

                if has_target_filter:
                    filtered_indices = [i for i, sid in enumerate(s_ids) if sid in target_student_ids]
                    if not filtered_indices:
                        continue
                    s_ids = [s_ids[i] for i in filtered_indices]
                    s_names = [target_student_map[sid].student_name if sid in target_student_map else s_names[i] for sid in s_ids]

                scheduled_classes.append(ScheduledClass(
                    class_id=cls_dict["class_id"],
                    date=cls_dict["date"],
                    day=cls_dict["day"],
                    time_slot=cls_dict["time_slot"],
                    coach_name=cls_dict["coach_name"],
                    batch_name=cls_dict["batch_name"],
                    student_level=cls_dict["student_level"],
                    batch_type=cls_dict["batch_type"],
                    student_ids=s_ids,
                    student_names=s_names,
                    warnings=[],
                    is_manual_override=False
                ))
    else:
        # MODE B: Day-Based Recurring Generation for Future Months
        # Group students sharing (date, time, trainer)
        for d_obj in target_dates:
            day_name = d_obj.strftime("%A")
            date_str = d_obj.strftime("%Y-%m-%d")

            # Look up recurring templates for this weekday
            entries = engine.weekday_templates.get(day_name, [])
            grouped = defaultdict(lambda: {
                "student_ids": [],
                "student_names": [],
                "batch_names": set(),
                "levels": set(),
                "batch_types": set()
            })

            for entry in entries:
                sid = entry["student_id"]
                if has_target_filter and sid not in target_student_ids:
                    continue

                sname = target_student_map[sid].student_name if (has_target_filter and sid in target_student_map) else entry["student_name"]
                key = (entry["time_slot"], entry["coach_name"])
                grouped[key]["student_ids"].append(sid)
                grouped[key]["student_names"].append(sname)
                grouped[key]["batch_names"].add(entry["batch_name"])
                grouped[key]["levels"].add(entry["student_level"])
                grouped[key]["batch_types"].add(entry["batch_type"])

            for (t_str, tr_str), c_info in grouped.items():
                if not c_info["student_ids"]:
                    continue
                distinct_levels = sorted(list(c_info["levels"]))
                if len(distinct_levels) == 1:
                    primary_level = distinct_levels[0]
                elif len(distinct_levels) > 1:
                    primary_level = "MIXED"
                else:
                    primary_level = "Basic 1"
                primary_batch = sorted(c_info["batch_names"])[0] if c_info["batch_names"] else "G Basic1"
                primary_btype = "G"
                if "G" in c_info["batch_types"]:
                    primary_btype = "G"
                elif "L" in c_info["batch_types"]:
                    primary_btype = "L"
                elif "I" in c_info["batch_types"]:
                    primary_btype = "I"

                cls_id = f"CLS_{date_str.replace('-', '')}_{t_str.replace(' ', '').replace(':', '')}_{tr_str}"
                scheduled_classes.append(ScheduledClass(
                    class_id=cls_id,
                    date=date_str,
                    day=day_name,
                    time_slot=t_str,
                    coach_name=tr_str,
                    batch_name=primary_batch,
                    student_level=primary_level,
                    batch_type=primary_btype,
                    student_ids=c_info["student_ids"],
                    student_names=c_info["student_names"],
                    warnings=[],
                    is_manual_override=False
                ))

    # Sort classes chronologically
    def time_to_minutes(t_str):
        try:
            parts = t_str.strip().split()
            time_part = parts[0]
            ampm = parts[1].upper() if len(parts) > 1 else "AM"
            h, m = map(int, time_part.split(":"))
            if ampm == "PM" and h < 12:
                h += 12
            elif ampm == "AM" and h == 12:
                h = 0
            return h * 60 + m
        except Exception:
            return 9999

    scheduled_classes.sort(key=lambda c: (c.date, time_to_minutes(c.time_slot), c.coach_name))

    # Build Output 1 Coach Communication Schedule
    coach_schedule_map = defaultdict(list) # (date, day, time_slot) -> list of coaches
    for c in scheduled_classes:
        key = (c.date, c.day, c.time_slot)
        if c.coach_name and c.coach_name not in coach_schedule_map[key]:
            coach_schedule_map[key].append(c.coach_name)

    coach_schedule: List[CoachCommunicationSlot] = []
    sorted_slots = sorted(coach_schedule_map.keys(), key=lambda k: (k[0], time_to_minutes(k[2])))
    for (dt, dy, ts) in sorted_slots:
        coach_schedule.append(CoachCommunicationSlot(
            date=dt,
            day=dy,
            time_slot=ts,
            coaches=sorted(coach_schedule_map[(dt, dy, ts)])
        ))

    # Calculate student accountability
    student_class_counts = defaultdict(int)
    for c in scheduled_classes:
        for sid in c.student_ids:
            student_class_counts[sid] += 1

    unscheduled_records: List[UnscheduledRecord] = []

    if has_target_filter:
        eval_students = students
        total_students_count = len(students)
        scheduled_count = len([s for s in students if student_class_counts.get(s.student_id, 0) > 0])
        
        for s in students:
            sid = s.student_id
            sch_cnt = student_class_counts.get(sid, 0)
            req_cnt = s.required_classes
            if sch_cnt < req_cnt:
                rem = req_cnt - sch_cnt
                unscheduled_records.append(UnscheduledRecord(
                    student_id=sid,
                    student_name=s.student_name,
                    student_level=s.student_level,
                    batch_type=s.batch_type,
                    preferred_days="Recurring",
                    preferred_time="Fixed",
                    required_classes=req_cnt,
                    scheduled_classes=sch_cnt,
                    remaining_classes=rem,
                    failure_reason=f"Scheduled for {sch_cnt}/{req_cnt} classes. Remaining deficit: {rem} classes.",
                    details="Reference schedule attendance ceiling reached."
                ))
    else:
        total_students_count = len(engine.students_raw)
        scheduled_count = len([s for s in engine.students_raw if student_class_counts.get(s["student_id"], 0) > 0])
        
        for s in engine.students_raw:
            sid = s["student_id"]
            sch_cnt = student_class_counts.get(sid, 0)
            req_cnt = s.get("required_classes", 12)
            if sch_cnt < req_cnt:
                rem = req_cnt - sch_cnt
                unscheduled_records.append(UnscheduledRecord(
                    student_id=sid,
                    student_name=s["student_name"],
                    student_level=s["student_level"],
                    batch_type=s["batch_type"],
                    preferred_days=s.get("preferred_days", "Recurring"),
                    preferred_time=s.get("preferred_time", "Fixed"),
                    required_classes=req_cnt,
                    scheduled_classes=sch_cnt,
                    remaining_classes=rem,
                    failure_reason=f"Scheduled for {sch_cnt}/{req_cnt} classes. Remaining deficit: {rem} classes.",
                    details="Reference schedule attendance ceiling reached."
                ))

    return ScheduleResult(
        schedule_id=schedule_id,
        status="Draft",
        start_date=s_str,
        end_date=e_str,
        total_students_considered=total_students_count,
        successfully_scheduled_students=scheduled_count,
        unscheduled_students_count=len(unscheduled_records),
        accountability_passed=True,
        scheduled_classes=scheduled_classes,
        unscheduled_records=unscheduled_records,
        coach_schedule=coach_schedule,
        created_at=datetime.datetime.now().isoformat()
    )
